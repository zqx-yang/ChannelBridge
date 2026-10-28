#!/usr/bin/env python3
"""Our best pose against iMARS's best linker fusion, per enzyme pair.

iMARS Table S2G is their own rank-1 prediction for each pair: the best linker
of their library, with the D_IC, O_IC and DO of the fusion AlphaFold2 built.
Our `poses.tsv` holds the best geometry a bispecific binder could hold. Both
are DO on the same coefficients, so they are directly comparable -- but they
are answers to different questions, and the difference is the point:

* **iMARS fuses the two enzymes N-to-C with a peptide linker.** The geometry
  is whatever that construct folds into. Their D_IC across these pairs runs
  10-55 A, because a terminal fusion cannot choose where the two active-site
  channels end up pointing.
* **We search for a pose a binder could hold.** D_IC is then limited by the
  two surfaces, not by the chain termini, so it can be much shorter -- and DO
  falls 0.0079 per angstrom of D_IC, so a shorter reachable D_IC raises DO
  mechanically.

So this is not a like-for-like contest, and a higher DO here is not a claim
that our construct works better. It is a measurement of how much geometric
room the binder route has that the fusion route does not. Two further caveats
belong with any use of these numbers:

* every enzyme here is now represented by coordinates of its OWN protein (3
  deposited entries, 2 switched to their own entry, 10 AlphaFill, 5 AlphaFold
  DB, 4 AlphaFold3), but 19 of the 24 are models, while theirs are AlphaFold2
  models of the actual fusion proteins. Two of the 24 are on an entry
  ChannelsDB annotated, and neither seeded run recovered that annotation, so
  those two tunnels are CAVER's own choice.
* a feasible pose is a geometry a binder **could** occupy. No binder has been
  designed, and iMARS measured yields; we have measured nothing.

The workbook is copyrighted and is not in this repository. Point
CHANNELBRIDGE_IMARS_ST2 at your own download of the supplement
(1-s2.0-S0092867424014715-mmc3.xlsx):

    export CHANNELBRIDGE_IMARS_ST2=/path/to/mmc3.xlsx
    pip install openpyxl
    ./compare_to_imars.py
"""

from __future__ import annotations

import csv
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

from run_examples import duplicate_pairs, pairs  # noqa: E402

#: Our pair -> the "Enzyme" cell of Table S2G. Their spellings vary ("4CL and
#: STS" vs "GPPS-MYS"), so the mapping is written out rather than guessed; a
#: fuzzy match here would quietly pair the wrong rows.
ST2G_OF = {
    "iMARS_01_4CL-STS": "4CL and STS",
    "iMARS_02_GPPS-MYS": "GPPS-MYS",
    "iMARS_03_GPPS-LIC": "GPPS-LIC",
    "iMARS_04_GPPS-OCS": "GPPS-OCS",
    "iMARS_05_4pOMT-7OMT": "4'OMT and 7OMT",
    "iMARS_06_4CL-BAS": "4CL and BAS",
    "iMARS_07_4pOMT-REPI": "4'OMT and REPI",
    "iMARS_08_GGPPS-TS": "GGPPS-TS",
    "iMARS_09_ERG20-NES": "ERG20-NES",
    "iMARS_10_ERG20-SQS": "ERG20-SQS",
    "iMARS_11_ERG20-ADS": "ERG20-ADS",
    "iMARS_12_4CL-RAS": "4CL and RAS",
    "iMARS_13_ERG20-GGPPS": "ERG20-GGPPS",
    "iMARS_14_GGPPS-AS": "GGPPS-AS",
    "iMARS_15_4CL-CHS": "4CL and CHS",
    "iMARS_16_GPPS-GES": "GPPS-GES",
    "iMARS_17_4pOMT-RNMT": "4'OMT and RNMT",
    "iMARS_18_4CL-CUS": "4CL and CUS",
    "iMARS_19_ERG20-BBS": "ERG20-BBS",
}


def read_st2g(path):
    """``{enzyme cell: (linker, D_IC, O_IC, DO)}`` from Table S2G."""
    try:
        import openpyxl
    except ImportError:
        raise SystemExit("this needs openpyxl: pip install openpyxl") from None
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "ST2G" not in wb.sheetnames:
        raise SystemExit("%s has no ST2G sheet (found %s). Table S2G is in "
                         "supplementary file mmc3.xlsx."
                         % (os.path.basename(path), ", ".join(wb.sheetnames)))
    out = {}
    seen_header = False
    for row in wb["ST2G"].iter_rows(values_only=True):
        if row and row[0] == "No.":
            seen_header = True
            continue
        if not seen_header or not row or row[0] is None:
            continue
        try:
            int(str(row[0]).strip())
        except ValueError:
            continue
        out[str(row[1]).strip()] = (str(row[5]).strip(), float(row[7]),
                                    float(row[8]), float(row[9]))
    wb.close()
    return out


def our_best(prefix):
    """``(D_IC, O_IC, DO, n_feasible)`` from a pair's poses.tsv, or None."""
    path = os.path.join(HERE, prefix, "output", "%s_poses.tsv" % prefix)
    if not os.path.exists(path):
        return None
    with open(path, encoding="utf-8") as fh:
        rows = [r for r in csv.DictReader(fh, delimiter="\t")
                if r["status"] == "ok"]
    if not rows:
        return None
    best = max(rows, key=lambda r: float(r["DO_clamped"]))
    return (float(best["D_IC"]), float(best["O_IC"]),
            float(best["DO_clamped"]), len(rows))


def main():
    path = os.environ.get("CHANNELBRIDGE_IMARS_ST2", "")
    if not path or not os.path.isfile(path):
        raise SystemExit(
            "set CHANNELBRIDGE_IMARS_ST2 to your own copy of the iMARS "
            "supplement mmc3.xlsx (it is copyrighted and not redistributed "
            "here); see the docstring.")
    theirs = read_st2g(path)
    # Two pairs built from one set of coordinates are one computation, not
    # two results, so they must not both enter an average. On this set there
    # are none -- but the check that said there were four was keyed on a
    # superseded table and silently removed three real pairs from these
    # statistics, so the question is asked of `inputs.tsv` and asked here.
    dup = {second for _first, second in duplicate_pairs()}

    print("%-22s %-11s %-11s %-11s %-9s %s"
          % ("pair", "our DO", "iMARS DO", "delta DO", "our D_IC", "their D_IC"))
    deltas, rows = [], []
    for p in pairs():
        prefix = p["prefix"]
        key = ST2G_OF.get(prefix)
        t = theirs.get(key) if key else None
        mine = our_best(prefix)
        note = "  (duplicate of an earlier pair)" if prefix in dup else ""
        if t is None:
            print("%-22s %-11s %-11s %-11s %-9s %s"
                  % (prefix, "-", "not in S2G", "-", "-", "-"))
            continue
        if mine is None:
            print("%-22s %-11s %-11.4f %-11s %-9s %.1f   no pose (no tunnel)%s"
                  % (prefix, "-", t[3], "-", "-", t[1], note))
            continue
        delta = mine[2] - t[3]
        print("%-22s %-11.4f %-11.4f %+-11.4f %-9.1f %.1f%s"
              % (prefix, mine[2], t[3], delta, mine[0], t[1], note))
        rows.append((prefix, mine, t, delta))
        if prefix not in dup:
            deltas.append(delta)

    if not deltas:
        return 1
    deltas.sort()
    n = len(deltas)
    median = deltas[n // 2] if n % 2 else 0.5 * (deltas[n // 2 - 1] + deltas[n // 2])
    print("\nover %d DISTINCT pair(s)%s:"
          % (n, " -- %d duplicate(s) excluded from the statistics" % len(dup)
             if dup else " (no two pairs are the same computation)"))
    print("  delta DO   min %+.4f   median %+.4f   max %+.4f"
          % (deltas[0], median, deltas[-1]))
    print("  our DO is higher for %d of %d" % (sum(1 for d in deltas if d > 0), n))
    ours_d = [r[1][0] for r in rows if r[0] not in dup]
    theirs_d = [r[2][1] for r in rows if r[0] not in dup]
    print("  D_IC       ours %.1f-%.1f A   theirs %.1f-%.1f A"
          % (min(ours_d), max(ours_d), min(theirs_d), max(theirs_d)))
    print("\nEach angstrom of D_IC is worth 0.0079 DO, so the gap is almost")
    print("all D_IC: a terminal fusion cannot place the exits where a binder can.")
    print("See the module docstring before quoting any of this.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
